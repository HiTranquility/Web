import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def create_report():
    doc = Document()

    # Set page margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Base directory for images
    img_dir = r"C:\Users\Admin\.gemini\antigravity-ide\brain\14870ee8-7f36-46e3-9412-1a13f25a116a"

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("BÁO CÁO NÂNG CẤP HỆ THỐNG VÀ TÍNH NĂNG MỚI")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(30, 58, 138) # Deep navy

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(24)
    run_sub = p_sub.add_run("Đồ án Web Đọc Truyện — Java Servlet & JSP (MVC Model 2)")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(100, 116, 139)

    # Metadata box
    table_meta = doc.add_table(rows=2, cols=2)
    table_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Nội dung thực hiện:", "4 Phân hệ Nâng cấp & Tính năng mới"),
        ("Công nghệ sử dụng:", "Java Servlet 3.1, JSP 2.3, MySQL, HikariCP, Vanilla CSS"),
    ]
    for row_idx, (k, v) in enumerate(meta_data):
        row = table_meta.rows[row_idx]
        cell_k, cell_v = row.cells[0], row.cells[1]
        cell_k.text = k
        cell_k.paragraphs[0].runs[0].font.bold = True
        cell_k.paragraphs[0].runs[0].font.size = Pt(10)
        cell_v.text = v
        cell_v.paragraphs[0].runs[0].font.size = Pt(10)
        set_cell_background(cell_k, "F1F5F9")
        set_cell_background(cell_v, "F8FAFC")
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Helper function to add headings
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = RGBColor(217, 119, 6) # Amber / Gold
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 41, 59)
        return p

    def add_body(text, bold_prefix=None, italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.2
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.name = "Arial"
            r_bold.font.size = Pt(10.5)
            r_bold.font.bold = True
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10.5)
        r.font.italic = italic
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.name = "Arial"
            r_bold.font.size = Pt(10.5)
            r_bold.font.bold = True
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10.5)
        return p

    def add_image_box(filename, caption):
        full_path = os.path.join(img_dir, filename)
        if os.path.exists(full_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(8)
            p_img.paragraph_format.space_after = Pt(2)
            run = p_img.add_run()
            run.add_picture(full_path, width=Inches(5.8))
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(12)
            r_cap = p_cap.add_run(f"Hình ảnh minh họa: {caption}")
            r_cap.font.name = "Arial"
            r_cap.font.size = Pt(9.5)
            r_cap.font.italic = True
            r_cap.font.color.rgb = RGBColor(100, 116, 139)
        else:
            add_body(f"[Ảnh không tìm thấy: {filename}]", italic=True)

    def add_sequence_box(steps_list):
        p_seq = doc.add_paragraph()
        p_seq.paragraph_format.space_before = Pt(6)
        p_seq.paragraph_format.space_after = Pt(6)
        table = doc.add_table(rows=len(steps_list) + 1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # Header
        hdr = table.rows[0]
        hdr.cells[0].text = "Bước"
        hdr.cells[1].text = "Tác tử (From → To)"
        hdr.cells[2].text = "Hành động & Dữ liệu xử lý"
        for c in hdr.cells:
            c.paragraphs[0].runs[0].font.bold = True
            c.paragraphs[0].runs[0].font.size = Pt(9.5)
            set_cell_background(c, "E2E8F0")
            
        for idx, (step_num, agents, desc) in enumerate(steps_list):
            row = table.rows[idx + 1]
            row.cells[0].text = str(step_num)
            row.cells[1].text = agents
            row.cells[2].text = desc
            for c in row.cells:
                c.paragraphs[0].runs[0].font.size = Pt(9)
                set_cell_background(c, "F8FAFC" if idx % 2 == 0 else "FFFFFF")
        doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # =========================================================================
    # PHẦN 1
    # =========================================================================
    add_h1("1. Nâng cấp Hệ thống Ví: Mở khóa chương VIP bằng Xu (ISSUE-020)")
    
    add_h2("1.1. Hiện trạng trước khi nâng cấp (Before)")
    add_bullet("Đã có bảng wallets và transactions từ giai đoạn trước (ISSUE-008).")
    add_bullet("Xu ảo chỉ dùng để gửi tặng (tip) tác giả một chiều. Người đọc không có động lực kiếm hoặc nạp xu, tác giả cũng không có cơ chế thu phí bản quyền nội dung.")
    add_bullet("Mọi chương truyện khi đọc tại /chapter?action=read đều mở tự do, không có cơ chế phân loại VIP hay kiểm soát truy cập.")

    add_image_box("before_chapter_read_view_1790818251072.png", "Giao diện đọc truyện trước khi có tính năng chương VIP (Đọc tự do)")

    add_h2("1.2. Tính năng nâng cấp mới (After)")
    add_bullet("Khi tác giả tạo/sửa chương có thể tick chọn 'Chương VIP' và định giá mở khoá bằng Xu (VD: 10, 20 xu).", "• Chương VIP (Paywall): ")
    add_bullet("Máy chủ chủ động làm rỗng văn bản (chapter.setContent(\"\")) và chặn endpoint xem thô (action=raw). Người dùng bấm F12 hoặc Ctrl+U hoàn toàn không thấy nội dung.", "• Bảo mật 100%: ")
    add_bullet("Kiểm tra số dư, trừ xu độc giả, cộng xu tác giả và ghi nhật ký giao dịch trong một transaction nguyên tử (con.setAutoCommit(false)) với khóa bi quan FOR UPDATE.", "• Giao dịch ACID an toàn: ")
    add_bullet("Sau khi mở khoá thành công, bản ghi được lưu vào bảng chapter_unlocks. Người đọc sở hữu vĩnh viễn, các lần đọc sau hoàn toàn miễn phí.", "• Quyền đọc vĩnh viễn: ")

    add_image_box("after_vip_chapter_locked_1790819619494.png", "Màn hình Paywall chặn chương VIP khi chưa mở khoá (Báo giá 20 xu & Nút mở khoá)")
    add_image_box("after_vip_chapter_unlocked_1790819646410.png", "Nội dung chương VIP được mở và hiển thị đầy đủ sau khi mở khoá")

    add_h2("1.3. Thay đổi kỹ thuật (Technical Changes)")
    add_bullet("Thêm cột is_vip TINYINT(1) DEFAULT 0, coin_price INT DEFAULT 0 vào bảng chapters; tạo bảng chapter_unlocks (user_id, chapter_id, coin_price, unlocked_at).", "• CSDL: ")
    add_bullet("Bổ sung thuộc tính vip, coinPrice trong Chapter.java; tạo model ChapterUnlock.java; viết UnlockDAO.java (hasUnlocked, findUnlockedChapterIds); thêm hàm transaction unlockChapter() trong WalletDAO.java.", "• Model & DAO: ")
    add_bullet("ChapterServlet.java kiểm tra paywall ở read() và xử lý POST action=unlock; tạo view _locked.jsp và bổ sung tùy chọn VIP trong form.jsp.", "• Controller & View: ")

    add_h2("1.4. Luồng dữ liệu và Sơ đồ tuần tự (Sequence Diagram)")
    add_image_box("diagram_seq_vip.png", "Sơ đồ tuần tự: Luồng mở khóa chương VIP bằng Xu (ISSUE-020)")
    vip_steps = [
        (1, "Độc giả → ChapterServlet", "Gửi request đọc chương GET /chapter?action=read&id=61"),
        (2, "ChapterServlet → UnlockDAO", "Gọi hasUnlocked(userId, 61) để kiểm tra trạng thái mở khóa"),
        (3, "UnlockDAO → MySQL", "SELECT COUNT(*) FROM chapter_unlocks WHERE user_id=? AND chapter_id=?"),
        (4, "MySQL → UnlockDAO", "Trả về 0 (Chưa mở khóa)"),
        (5, "ChapterServlet → ChapterServlet", "chapter.setContent(\"\") để bảo mật tuyệt đối mã nguồn HTML"),
        (6, "ChapterServlet → JSP View", "Forward sang read.jsp với cờ isLocked=true và nạp _locked.jsp"),
        (7, "Độc giả → ChapterServlet", "Bấm 'Mở khoá ngay' (POST /chapter?action=unlock&id=61)"),
        (8, "ChapterServlet → WalletDAO", "Gọi unlockChapter(userId, authorId, chapterId, 20)"),
        (9, "WalletDAO → MySQL", "Mở Transaction, khóa SELECT balance FROM wallets FOR UPDATE"),
        (10, "WalletDAO → MySQL", "Trừ ví độc giả (-20), cộng ví tác giả (+20), ghi transactions & chapter_unlocks"),
        (11, "MySQL → WalletDAO", "Commit transaction thành công"),
        (12, "ChapterServlet → Độc giả", "Redirect về trang đọc, nạp lại nội dung chương hoàn chỉnh")
    ]
    add_sequence_box(vip_steps)

    # =========================================================================
    # PHẦN 2
    # =========================================================================
    add_h1("2. Nâng cấp Đánh giá: Review chuyên sâu + Spoiler Tag + Vote hữu ích (ISSUE-022)")
    
    add_h2("2.1. Hiện trạng trước khi nâng cấp (Before)")
    add_bullet("Chỉ có form chấm sao 1–5 sao đơn giản, không viết được nhận xét hay cảm nhận về tác phẩm.")
    add_bullet("Mục bình luận thường xuất hiện tình trạng tiết lộ trước tình tiết truyện (spoiler), gây ức chế cho người đọc mới.")
    add_bullet("Thanh điều hướng truyện chỉ có 2 tab cơ bản: Mục lục và Bình luận.")

    add_image_box("before_story_rating_area_1790818241601.png", "Trang chi tiết truyện trước khi có tab Đánh giá chuyên sâu (Chỉ có 2 tab)")

    add_h2("2.2. Tính năng nâng cấp mới (After)")
    add_bullet("Độc giả được viết Tiêu đề bài viết và Nội dung cảm nhận dài (phân tích cốt truyện, tâm lý nhân vật, văn phong).", "• Bài viết đánh giá chi tiết: ")
    add_bullet("Checkbox 'Đánh giá có chứa Spoiler'. Nội dung review bị làm mờ bằng CSS (filter: blur(6px)) và cấm bôi đen, bấm nút 'Hiển thị nội dung' mới giải mã.", "• Cảnh báo Spoiler thông minh: ")
    add_bullet("Độc giả khác bấm '👍 Có ích'. Các đánh giá nhận nhiều bình chọn nhất được ưu tiên nổi bật trên đầu tab.", "• Bình chọn hữu ích (Helpful Votes): ")

    add_image_box("after_story_toc_vip_badge_1790819731799.png", "Thanh điều hướng truyện tích hợp Tab thứ 3 'Đánh giá' và huy hiệu chương VIP")

    add_h2("2.3. Thay đổi kỹ thuật (Technical Changes)")
    add_bullet("Tạo bảng reviews (id, user_id, story_id, title, content, has_spoiler, helpful_count, status...) và bảng review_votes (user_id, review_id, voted_at).", "• CSDL: ")
    add_bullet("Tạo Review.java; xây dựng ReviewDAO.java với các hàm findByStory, findByUserAndStory, upsert (1 người/truyện) và toggleVote.", "• Model & DAO: ")
    add_bullet("RatingServlet.java hỗ trợ action=review và action=vote; StoryServlet.java nạp danh sách reviews; tạo mảnh _review.jsp và mở rộng detail.jsp.", "• Controller & View: ")

    add_h2("2.4. Luồng dữ liệu và Sơ đồ tuần tự (Sequence Diagram)")
    add_image_box("diagram_seq_review.png", "Sơ đồ tuần tự: Luồng đánh giá chi tiết, Spoiler Tag & Vote hữu ích")
    review_steps = [
        (1, "Độc giả → RatingServlet", "Gửi đánh giá: POST /rating?action=review (title, content, hasSpoiler=1)"),
        (2, "RatingServlet → ReviewDAO", "Gọi upsert(review)"),
        (3, "ReviewDAO → MySQL", "INSERT INTO reviews ... ON DUPLICATE KEY UPDATE title=..., content=..."),
        (4, "RatingServlet → Độc giả", "Redirect về trang truyện /story?action=detail&id=1#reviews"),
        (5, "Độc giả → StoryServlet", "Tải trang truyện GET /story?action=detail&id=1#reviews"),
        (6, "StoryServlet → ReviewDAO", "Gọi findByStory(storyId, currentUserId)"),
        (7, "ReviewDAO → MySQL", "Truy vấn danh sách review kèm số lượt vote và cờ voted_by_me"),
        (8, "StoryServlet → JSP View", "Render detail.jsp và _review.jsp (Thẻ bài viết có class .has-spoiler)"),
        (9, "Độc giả → Trình duyệt", "Bấm nút 'Hiển thị nội dung' → JS gỡ bỏ lớp mờ tại máy khách tức thì"),
        (10, "Độc giả → RatingServlet", "Bấm '👍 Có ích' (POST /rating?action=vote&reviewId=...)"),
        (11, "RatingServlet → ReviewDAO", "Gọi toggleVote(userId, reviewId) để tăng/giảm điểm và lưu lượt vote")
    ]
    add_sequence_box(review_steps)

    # =========================================================================
    # PHẦN 3
    # =========================================================================
    add_h1("3. Gamification: Điểm danh nhận Xu & Nhiệm vụ hàng ngày (ISSUE-026)")
    
    add_h2("3.1. Hiện trạng trước khi nâng cấp (Before)")
    add_bullet("Trang cá nhân /user?action=me chỉ hiển thị số dư tĩnh Ví xu ảo: 100 xu.")
    add_bullet("Không có cơ chế tương tác hàng ngày để giữ chân người đọc (retention) và tạo thói quen quay lại website.")
    add_bullet("Độc giả không có nguồn xu miễn phí để trải nghiệm đọc thử các chương trả phí.")

    add_image_box("before_user_wallet_profile_1790818214190.png", "Trang cá nhân trước khi nâng cấp Gamification (Chỉ có thông tin ví tĩnh)")

    add_h2("3.2. Tính năng nâng cấp mới (After)")
    add_bullet("Lưới chuỗi 7 ngày liên tục với phần thưởng tăng tiến (+5, +10, +15, +20, +25, +30, +50 xu). Đứt chuỗi sẽ tự động tính lại từ ngày 1.", "• Chuỗi điểm danh 7 ngày: ")
    add_bullet("Hệ thống tự động theo dõi hành vi: Điểm danh ngày mới (+10 xu), Mọt sách siêng năng đọc 1 chương (+10 xu), Nhà phê bình bình luận/đánh giá (+15 xu).", "• Nhiệm vụ độc giả hàng ngày: ")
    add_bullet("Kiểm tra điều kiện trực tiếp trên CSDL (view_logs, comments, reviews) và cộng xu nguyên tử vào ví người dùng.", "• Nhận thưởng an toàn: ")

    add_image_box("after_user_gamification_profile_1790819502803.png", "Phân hệ Trạm Nhận Xu với Chuỗi 7 ngày và Bảng nhiệm vụ hàng ngày")

    add_h2("3.3. Thay đổi kỹ thuật (Technical Changes)")
    add_bullet("Tạo bảng daily_checkins (id, user_id, checkin_date, streak_days, reward_coins) và daily_quest_claims (user_id, quest_key, claim_date, reward_coins).", "• CSDL: ")
    add_bullet("Tạo DailyCheckin.java, DailyQuest.java; xây dựng GamificationDAO.java với các hàm doCheckin, getDailyQuests, claimQuestReward.", "• Model & DAO: ")
    add_bullet("UserServlet.java xử lý action=checkin và action=claim-quest; bổ sung khối giao diện Gamification Dashboard hiện đại trong me.jsp.", "• Controller & View: ")

    add_h2("3.4. Luồng dữ liệu và Sơ đồ tuần tự (Sequence Diagram)")
    add_image_box("diagram_seq_gamification.png", "Sơ đồ tuần tự: Gamification điểm danh 7 ngày & Nhiệm vụ hàng ngày")
    quest_steps = [
        (1, "Độc giả → UserServlet", "Truy cập hồ sơ cá nhân GET /user?action=me"),
        (2, "UserServlet → GamificationDAO", "Gọi getTodayCheckin() và getDailyQuests()"),
        (3, "GamificationDAO → MySQL", "Kiểm tra chuỗi streak dựa trên checkin_date hôm qua; quét view_logs và comments"),
        (4, "UserServlet → JSP View", "Render me.jsp hiển thị chuỗi 7 ngày và trạng thái tiến độ các nhiệm vụ"),
        (5, "Độc giả → UserServlet", "Bấm '✨ Điểm Danh Ngay' (POST /user?action=checkin)"),
        (6, "UserServlet → GamificationDAO", "Gọi doCheckin(userId)"),
        (7, "GamificationDAO → MySQL", "Mở Transaction: Ghi daily_checkins, cộng xu wallets, ghi transactions"),
        (8, "MySQL → GamificationDAO", "Commit thành công, trả về đối tượng DailyCheckin"),
        (9, "UserServlet → Độc giả", "Redirect về trang cá nhân, thông báo nhận xu thành công và cập nhật số dư")
    ]
    add_sequence_box(quest_steps)

    # =========================================================================
    # PHẦN 4
    # =========================================================================
    add_h1("4. Trang chủ & Bảng xếp hạng: Vinh danh Tác giả & Xếp hạng Tương tác")
    
    add_h2("4.1. Hiện trạng trước khi nâng cấp (Before)")
    add_bullet("Trang chủ sơ khai, chỉ hiển thị danh sách truyện cơ bản, thiếu điểm nhấn vinh danh cộng đồng và lối tắt đến trạm kiếm xu.")
    add_bullet("Trang xếp hạng chỉ đo lượt xem truyện thụ động, hoàn toàn thiếu Bảng xếp hạng Tác giả nổi bật để khích lệ người sáng tác.")

    add_image_box("before_homepage_full_1790818258039.png", "Giao diện trang chủ ban đầu trước khi nâng cấp")

    add_h2("4.2. Tính năng nâng cấp mới (After)")
    add_bullet("Thuật toán xếp hạng tác giả dựa trên tổng lượt xem toàn bộ tác phẩm (SUM(stories.view_count)) và lượng fan theo dõi (users.follower_count).", "• Bảng vinh danh Tác giả: ")
    add_bullet("Huy hiệu Top 1, Top 2, Top 3 mạ kim loại (Vàng, Bạc, Đồng) trang trọng, có avatar và nút Xem hồ sơ tác giả.", "• Huy hiệu kim loại Top 1-2-3: ")
    add_bullet("Hệ thống tab gồm 7 tiêu chí: Tuần này, Tháng này, Xem nhiều nhất, Điểm cao, Nhiều chương, Mới đăng, Tác giả nổi bật.", "• Bộ lọc đa chiều: ")
    add_bullet("Tái cấu trúc CSS Flexbox cho mục lục popup đọc truyện (layout-reader.css và components.css), khắc phục hoàn toàn hiện tượng văng nhãn VIP lên góc tiêu đề.", "• Tinh chỉnh giao diện Mục lục: ")

    add_image_box("after_rank_gamification_1790819533062.png", "Bảng xếp hạng Tác giả nổi bật với huy hiệu thứ hạng và số liệu chi tiết")
    add_image_box("fixed_reader_toc_popup_1790820790070.png", "Mục lục popup đọc truyện sau khi sửa triệt để lỗi đè nhãn VIP (Căn thẳng hàng chuẩn xác)")

    add_h2("4.3. Thay đổi kỹ thuật (Technical Changes)")
    add_bullet("Bổ sung truy vấn tổng hợp SQL phức hợp (JOIN stories, LEFT JOIN user_follows, GROUP BY u.id ORDER BY total_views DESC) trong UserDAO.findTopAuthors().", "• DAO Layer: ")
    add_bullet("RankServlet.java mở rộng tham số lọc type=authors; rank/index.jsp render giao diện thẻ tác giả với avatar và thông số chi tiết.", "• Controller & View: ")
    add_bullet("Sửa triệt để lỗi position: absolute của class .badge thành .badge-vip với position: static !important trong components.css và layout-reader.css.", "• CSS Fix: ")

    add_h2("4.4. Luồng dữ liệu và Sơ đồ tuần tự (Sequence Diagram)")
    add_image_box("diagram_seq_rank.png", "Sơ đồ tuần tự: Bảng xếp hạng Tác giả nổi bật & Vinh danh")
    rank_steps = [
        (1, "Người dùng → RankServlet", "Bấm tab 'Tác giả nổi bật' GET /rank?type=authors"),
        (2, "RankServlet → UserDAO", "Gọi findTopAuthors(limit = 20)"),
        (3, "UserDAO → MySQL", "Thực thi SQL tổng hợp views từ stories và follows từ user_follows nhóm theo tác giả"),
        (4, "MySQL → UserDAO", "Trả về ResultSet danh sách 20 tác giả hàng đầu"),
        (5, "UserDAO → RankServlet", "Trả về List<User> topAuthors"),
        (6, "RankServlet → JSP View", "Forward sang rank/index.jsp với danh sách tác giả và huy hiệu Top 1-2-3"),
        (7, "JSP View → Người dùng", "Hiển thị bảng vinh danh trang trọng kèm nút xem hồ sơ từng tác giả")
    ]
    add_sequence_box(rank_steps)

    # Save to workspace
    out_path = r"c:\Users\Admin\Downloads\Web\DoAn-WebDocTruyen\BaoCao_NangCap_WebDocTruyen.docx"
    doc.save(out_path)
    print("SUCCESS: Word report generated at: " + out_path)

if __name__ == "__main__":
    create_report()
